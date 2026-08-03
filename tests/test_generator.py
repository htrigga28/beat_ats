from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from generator import generate_resume_docx
from schemas import ResumeDocument


def sample_resume() -> ResumeDocument:
    return ResumeDocument.model_validate(
        {
            "contact": {
                "full_name": "Jane Doe",
                "email": "jane@example.com",
                "phone": "+1 555 0100",
                "location": "Remote",
                "links": ["linkedin.com/in/janedoe"],
            },
            "professional_summary": "Frontend engineer focused on accessible products.",
            "work_experience": [
                {
                    "employer": "Example Labs",
                    "title": "Frontend Engineer",
                    "location": "Remote",
                    "start_date": "2022",
                    "end_date": "Present",
                    "bullets": [
                        {"text": "Built accessible React interfaces."},
                        {"text": "Partnered with product and design teams."},
                    ],
                }
            ],
            "skills": [{"label": "Frontend", "items": ["React", "TypeScript", "WCAG"]}],
            "education": [
                {
                    "institution": "Example University",
                    "credential": "BSc Computer Science",
                    "dates": "2018",
                }
            ],
            "certifications": ["Accessibility Fundamentals"],
            "projects": [
                {
                    "name": "Design System",
                    "link": "example.com/design-system",
                    "bullets": [{"text": "Created reusable UI components."}],
                }
            ],
        }
    )


def test_generator_creates_parseable_single_column_document() -> None:
    payload = generate_resume_docx(sample_resume())
    document = Document(BytesIO(payload))

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Jane Doe" in text
    assert "PROFESSIONAL SUMMARY" in text
    assert "WORK EXPERIENCE" in text
    assert "SKILLS" in text
    assert "EDUCATION" in text
    assert len(document.sections) == 1
    assert len(document.tables) == 0
    assert abs(document.sections[0].top_margin - Inches(0.72)) < 200
    assert abs(document.sections[0].left_margin - Inches(0.72)) < 200
    assert document.sections[0].header.paragraphs[0].text == ""
    assert document.sections[0].footer.paragraphs[0].text == ""


def test_generator_contains_no_ats_hostile_ooxml() -> None:
    payload = generate_resume_docx(sample_resume())

    with ZipFile(BytesIO(payload)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")

    assert "<w:tbl" not in document_xml
    assert "<w:txbxContent" not in document_xml
    assert "<w:drawing" not in document_xml
    assert "header" not in relationships.lower()
    assert "footer" not in relationships.lower()
    assert "FFFFFF" not in document_xml.upper()
    assert "<w:numPr" in document_xml


def test_generator_omits_empty_optional_sections_and_internal_ids() -> None:
    resume = ResumeDocument.model_validate({"contact": {"full_name": "Jane Doe"}})

    document = Document(BytesIO(generate_resume_docx(resume)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "WORK EXPERIENCE" not in text
    assert "SKILLS" not in text
    assert resume.contact.full_name in text


def test_generator_matches_preview_typography_and_entry_structure() -> None:
    document = Document(BytesIO(generate_resume_docx(sample_resume())))

    name = document.paragraphs[0]
    assert name.text == "Jane Doe"
    assert name.runs[0].font.size == Pt(18)
    assert name.runs[0].font.color.rgb == RGBColor(15, 23, 42)
    assert name.runs[0].font.all_caps is True

    contact = document.paragraphs[1]
    assert contact.text == "jane@example.com | +1 555 0100 | Remote"
    assert contact.runs[0].font.size == Pt(8.5)

    section_heading = next(
        paragraph for paragraph in document.paragraphs if paragraph.text == "WORK EXPERIENCE"
    )
    borders = section_heading._p.pPr.find(qn("w:pBdr"))
    assert borders is not None
    assert borders.find(qn("w:bottom")).get(qn("w:color")) == "94A3B8"

    role_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Frontend Engineer")
    )
    assert role_heading.text == "Frontend Engineer\tExample Labs"
    assert all(run.bold for run in role_heading.runs)

    normal = document.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size == Pt(9.5)
    assert normal.font.color.rgb == RGBColor(51, 65, 85)
