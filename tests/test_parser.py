from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from parser import (
    CorruptDocumentError,
    DocumentTooLargeError,
    UnsupportedFileTypeError,
    clean_text,
    extract_docx_text,
    extract_pdf_text,
    extract_resume,
)


def _docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("\u2013 Built   accessible\u00a0interfaces", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "React"
    table.cell(0, 1).text = "TypeScript"
    document.add_paragraph("Education")
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes(text: str | None = None, pages: int = 1) -> bytes:
    stream = BytesIO()
    pdf = canvas.Canvas(stream)
    for index in range(pages):
        if text:
            pdf.drawString(72, 760, f"{text} page {index + 1}")
        pdf.showPage()
    pdf.save()
    return stream.getvalue()


def _tightly_kerned_pdf_bytes() -> bytes:
    stream = BytesIO()
    pdf = canvas.Canvas(stream)
    first = "Improved"
    pdf.drawString(72, 760, first)
    next_x = 72 + stringWidth(first, "Helvetica", 12) + 2.5
    pdf.drawString(next_x, 760, "frontend data handling across production interfaces")
    pdf.showPage()
    pdf.save()
    return stream.getvalue()


def test_clean_text_normalizes_unicode_bullets_and_whitespace() -> None:
    raw = "  ●  Built\u00a0  interfaces\r\n\r\n\r\n▪ Improved “accessibility”\u0000  "

    assert clean_text(raw) == '• Built interfaces\n\n• Improved "accessibility"'


def test_extract_docx_preserves_paragraph_and_table_order() -> None:
    result = extract_docx_text(_docx_bytes())

    assert result.file_type == "docx"
    assert result.needs_vision_fallback is False
    assert result.text.splitlines() == [
        "Jane Doe",
        "• Built accessible interfaces",
        "React | TypeScript",
        "Education",
    ]


def test_extract_pdf_marks_image_only_document_for_vision() -> None:
    result = extract_pdf_text(_pdf_bytes())

    assert result.page_count == 1
    assert result.text == ""
    assert result.needs_vision_fallback is True
    assert result.warnings


def test_extract_pdf_preserves_spaces_between_tightly_kerned_words() -> None:
    result = extract_pdf_text(_tightly_kerned_pdf_bytes())

    assert "Improved frontend data handling" in result.text


def test_extract_pdf_rejects_too_many_pages() -> None:
    with pytest.raises(DocumentTooLargeError, match="20 pages"):
        extract_pdf_text(_pdf_bytes("resume", pages=21))


def test_extract_resume_validates_extension_and_signature() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_resume(b"not a pdf", "resume.pdf")

    with pytest.raises(UnsupportedFileTypeError):
        extract_resume(_pdf_bytes("resume"), "resume.txt")


def test_extract_resume_rejects_oversized_upload_before_parsing() -> None:
    with pytest.raises(DocumentTooLargeError, match="10 MB"):
        extract_resume(b"%PDF-" + b"0" * (10 * 1024 * 1024), "resume.pdf")


def test_extract_docx_rejects_invalid_zip() -> None:
    with pytest.raises(CorruptDocumentError):
        extract_docx_text(b"PK\x03\x04not-a-docx")
