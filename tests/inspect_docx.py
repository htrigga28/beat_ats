"""Return the semantic DOCX facts used by the browser integration test."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def inspect_docx(path: Path) -> dict[str, object]:
    document = Document(path)
    return {
        "paragraphs": [paragraph.text for paragraph in document.paragraphs if paragraph.text],
        "table_count": len(document.tables),
        "header_text": [
            paragraph.text
            for section in document.sections
            for paragraph in section.header.paragraphs
            if paragraph.text
        ],
        "footer_text": [
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
            if paragraph.text
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    args = parser.parse_args()
    print(json.dumps(inspect_docx(args.document)))


if __name__ == "__main__":
    main()
