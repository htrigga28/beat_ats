"""Generate deterministic, single-column, ATS-friendly Word resumes."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from schemas import ResumeDocument

INK = RGBColor(15, 23, 42)
BODY = RGBColor(51, 65, 85)
MUTED = RGBColor(71, 85, 105)
RULE = "94A3B8"
FONT_NAME = "Arial"
CONTENT_WIDTH_INCHES = 7.06


def _set_run_font(
    run: Run,
    size: float,
    *,
    bold: bool = False,
    color: RGBColor = BODY,
    all_caps: bool = False,
) -> None:
    font = run.font
    font.name = FONT_NAME
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    font.all_caps = all_caps
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)


def _configure_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BODY
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.15

    styles = document.styles
    if "Resume Section" not in styles:
        style = styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Resume Section"]
    style.font.name = FONT_NAME
    style.font.size = Pt(10.5)
    style.font.bold = True
    style.font.color.rgb = INK
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    style.paragraph_format.space_before = Pt(11)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True


def _set_paragraph_bottom_border(
    paragraph: Paragraph,
    *,
    color: str,
    size: int,
    space: int,
) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)


def _add_numbering(document: DocumentObject) -> int:
    """Create a real Word bullet definition with compact resume geometry."""

    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "bullet")
    level.append(number_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)

    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "360")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "360")
    indentation.set(qn("w:hanging"), "180")
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)

    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    run_properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "334155")
    run_properties.append(color)
    level.append(run_properties)
    abstract.append(level)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return num_id


def _add_section_heading(document: DocumentObject, title: str) -> None:
    paragraph = document.add_paragraph(style="Resume Section")
    _set_run_font(paragraph.add_run(title.upper()), 10.5, bold=True, color=INK)
    _set_paragraph_bottom_border(paragraph, color=RULE, size=6, space=3)


def _add_bullet(document: DocumentObject, text: str, num_id: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.15
    properties = paragraph._p.get_or_add_pPr()
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    number_properties.append(level)
    number_properties.append(number)
    properties.append(number_properties)
    run = paragraph.add_run(text)
    _set_run_font(run, 9.5)


def _add_labeled_line(document: DocumentObject, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}: ")
    _set_run_font(label_run, 9.5, bold=True, color=INK)
    value_run = paragraph.add_run(value)
    _set_run_font(value_run, 9.5)


def _add_entry_heading(document: DocumentObject, primary: str, secondary: str | None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    _set_run_font(paragraph.add_run(primary), 9.5, bold=True, color=INK)
    if secondary:
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(CONTENT_WIDTH_INCHES),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.SPACES,
        )
        _set_run_font(paragraph.add_run(f"\t{secondary}"), 9.5, bold=True, color=INK)


def _add_metadata(document: DocumentObject, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    _set_run_font(paragraph.add_run(value), 8.5, color=MUTED)


def _pack_contact_links(links: list[str], *, max_chars: int = 92) -> list[str]:
    lines: list[str] = []
    current: list[str] = []
    current_length = 0
    for link in links:
        added_length = len(link) + (3 if current else 0)
        if current and current_length + added_length > max_chars:
            lines.append(" | ".join(current))
            current = [link]
            current_length = len(link)
        else:
            current.append(link)
            current_length += added_length
    if current:
        lines.append(" | ".join(current))
    return lines


def _add_contact_line(
    document: DocumentObject,
    value: str,
    *,
    after: float,
    add_rule: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    _set_run_font(paragraph.add_run(value), 8.5, color=MUTED)
    if add_rule:
        _set_paragraph_bottom_border(paragraph, color="334155", size=12, space=7)


def generate_resume_docx(resume: ResumeDocument) -> bytes:
    """Return a parseable Word document without ATS-hostile layout constructs."""

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    _configure_styles(document)
    bullet_num_id = _add_numbering(document)

    properties = document.core_properties
    properties.title = "Tailored Resume"
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.keywords = ""

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(3)
    _set_run_font(
        name.add_run(resume.contact.full_name),
        18,
        bold=True,
        color=INK,
        all_caps=True,
    )

    primary_contact_parts = [
        part
        for part in [
            resume.contact.email,
            resume.contact.phone,
            resume.contact.location,
        ]
        if part
    ]
    contact_lines = []
    if primary_contact_parts:
        contact_lines.append(" | ".join(primary_contact_parts))
    contact_lines.extend(_pack_contact_links(resume.contact.links))
    for index, contact_line in enumerate(contact_lines):
        is_last = index == len(contact_lines) - 1
        _add_contact_line(
            document,
            contact_line,
            after=9 if is_last else 1,
            add_rule=is_last,
        )
    if not contact_lines:
        _set_paragraph_bottom_border(name, color="334155", size=12, space=7)
        name.paragraph_format.space_after = Pt(9)

    if resume.professional_summary:
        _add_section_heading(document, "Professional Summary")
        paragraph = document.add_paragraph(resume.professional_summary)
        paragraph.paragraph_format.space_after = Pt(4)

    if resume.work_experience:
        _add_section_heading(document, "Work Experience")
        for role in resume.work_experience:
            _add_entry_heading(document, role.title, role.employer)

            metadata = " | ".join(
                part
                for part in [
                    role.location,
                    " - ".join(part for part in [role.start_date, role.end_date] if part),
                ]
                if part
            )
            if metadata:
                _add_metadata(document, metadata)
            for bullet in role.bullets:
                _add_bullet(document, bullet.text, bullet_num_id)

    if resume.skills:
        _add_section_heading(document, "Skills")
        for group in resume.skills:
            value = ", ".join(group.items)
            if group.label:
                _add_labeled_line(document, group.label, value)
            elif value:
                paragraph = document.add_paragraph(value)
                paragraph.paragraph_format.space_after = Pt(2)

    if resume.education:
        _add_section_heading(document, "Education")
        for education in resume.education:
            credential = education.credential
            if education.field_of_study:
                credential = f"{credential}, {education.field_of_study}"
            _add_entry_heading(document, credential, education.institution)
            metadata = " | ".join(part for part in [education.location, education.dates] if part)
            if metadata:
                _add_metadata(document, metadata)
            for detail in education.details:
                _add_bullet(document, detail, bullet_num_id)

    if resume.certifications:
        _add_section_heading(document, "Certifications")
        for certification in resume.certifications:
            _add_bullet(document, certification, bullet_num_id)

    if resume.projects:
        _add_section_heading(document, "Projects")
        for project in resume.projects:
            _add_entry_heading(document, project.name, project.role)
            metadata = " | ".join(part for part in [project.dates, project.link] if part)
            if metadata:
                _add_metadata(document, metadata)
            for bullet in project.bullets:
                _add_bullet(document, bullet.text, bullet_num_id)

    for additional in resume.additional_sections:
        if not additional.items:
            continue
        _add_section_heading(document, additional.title)
        for item in additional.items:
            _add_bullet(document, item, bullet_num_id)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
